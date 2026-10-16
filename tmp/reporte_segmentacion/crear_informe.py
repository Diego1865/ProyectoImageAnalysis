from pathlib import Path
from copy import deepcopy
import hashlib
import json
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
TMP = Path(__file__).parent
BASE = ROOT/'PracticaSegmentacion'
E = BASE/'evidencia'
REF = Path(r'C:\Users\dibao\Downloads\plantilla_informe_practica-Segmentación.docx')
OUT = BASE/'Practica Segmentación.docx'
D = json.loads((E/'resultados.json').read_text(encoding='utf-8'))
A = D['images']['monedas_separadas']
B = D['images']['monedas_juntas']
assert hashlib.sha256(REF.read_bytes()).hexdigest() == 'ae4b8d3a92b1d709dd93db71de7090b1f6c9ef45adc0c6aca996c5493f09378a'
doc = Document(REF)
table_proto = deepcopy(doc.tables[0]._element)
for el in list(doc._element.body):
    if el.tag != qn('w:sectPr'):
        doc._element.body.remove(el)
for name in ['Title','Subtitle','Heading 1','Heading 2']:
    st = doc.styles[name]
    st.font.color.rgb = RGBColor(0, 0, 0)
    for c in st._element.xpath('.//w:color'):
        for key in ['themeColor','themeTint','themeShade']:
            c.attrib.pop(qn('w:'+key), None)
    st.paragraph_format.keep_with_next = True


def p(text='', style=None):
    para = doc.add_paragraph(text, style)
    return para


def h(text, level=2):
    return p(text, 'Heading '+str(level))


def page(title=None, level=2):
    if title:
        para = h(title, level)
        para.paragraph_format.page_break_before = True


def figure(path, caption, width=6.5):
    para = p()
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.keep_with_next = True
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.add_run().add_picture(str(path), width=Inches(width))
    cap = p(caption)
    cap.paragraph_format.space_after = Pt(8)
    for r in cap.runs:
        r.font.size = Pt(9)
    return cap


def table(headers, rows, widths=None, from_template=False):
    if from_template:
        el = deepcopy(table_proto)
        doc._element.body.insert(len(doc._element.body)-1, el)
        t = doc.tables[-1]
        for row in list(t._element.tr_lst)[1:]:
            t._element.remove(row)
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = str(text)
    else:
        t = doc.add_table(rows=1, cols=len(headers))
        for c, text in zip(t.rows[0].cells, headers):
            c.text = str(text)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for values in rows:
        for cell, text in zip(t.add_row().cells, values):
            cell.text = str(text)
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    pr = t._element.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge = OxmlElement('w:'+side)
        for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]:
            edge.set(qn('w:'+key), val)
        borders.append(edge)
    old = pr.find(qn('w:tblBorders'))
    if old is not None:
        pr.remove(old)
    pr.append(borders)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            for old in list(tcpr.findall(qn('w:tcBorders'))):
                tcpr.remove(old)
            cell_borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge = OxmlElement('w:'+side)
                for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    edge.set(qn('w:'+key), val)
                cell_borders.append(edge)
            tcpr.append(cell_borders)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8E8E8' if i == 0 else 'FFFFFF')
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                m = OxmlElement('w:'+side)
                m.set(qn('w:w'), '75')
                m.set(qn('w:type'), 'dxa')
                margins.append(m)
            tcpr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.space_before = Pt(2)
                para.paragraph_format.line_spacing = 1
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.font.bold = i == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p().paragraph_format.space_after = Pt(0)
    return t


def threshold_row(data, stage, method):
    return next(r for r in data['thresholds'] if r['stage']==stage and r['method']==method)


p('Informe de la Práctica 4', 'Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
p('Segmentación de regiones y ajuste de brillo', 'Subtitle')
p('Integrantes: Barragán Osorio Diego José María & Flores Luna Eduardo Daniel')
p('Grupo / Comunidad: __________________________')
p('Fecha: 16 de octubre de 2026')
h('Introducción', 1)
p('La segmentación por intensidad permite separar regiones con niveles de gris diferentes. Su resultado depende de la iluminación, el contraste, el ruido y la proximidad entre objetos. En esta práctica se analizaron dos fotografías de una misma escena y se eligieron las tres monedas como objetos de interés. Una fotografía presenta las monedas separadas y la otra las muestra próximas entre sí.')
p('El experimento encontró tres regiones principales en la escena con monedas separadas. En la escena con monedas juntas, las dos monedas superiores quedaron conectadas y el procesamiento final produjo dos regiones. Este resultado muestra la necesidad de revisar la forma de la máscara además del conteo.')
h('Objetivo')
p('Aplicar y comparar técnicas de umbralado y ajuste de brillo mediante histogramas y funciones de potencia; evaluar el efecto del ruido y los filtros; y obtener área, perímetro y número de componentes conexas de las regiones segmentadas.')
h('Alcance del experimento')
p('Se compararon Otsu, Kapur, mínimo del histograma, media, multiumbralización y umbral por banda. Se ensayaron las transformaciones de brillo de la guía, junto con desplazamiento, contracción y expansión. El análisis se realizó sobre un recorte fijo que contiene las tres monedas, sin redimensionar las fotografías.')
p('La actividad se identifica como Práctica 4 conforme a la asignación; el PDF y la plantilla proporcionados están numerados como Práctica 5. La aplicación principal se complementó con un script de experimentación independiente para obtener las evidencias de este informe.')

page('Fundamentos y ejercicio a mano')
p('Para una imagen de 8 bits, el histograma h(i) cuenta los píxeles de intensidad i, con i entre 0 y 255. Su versión normalizada p(i)=h(i)/N representa frecuencias relativas. En estas fotografías las monedas son más oscuras que la mesa; por ello se asignó blanco al objeto cuando I≤T y negro al fondo cuando I>T.')
p('Otsu busca el umbral que minimiza la varianza dentro de las dos clases, equivalente a maximizar su separación entre clases. Kapur maximiza H₀(T)+H₁(T), donde cada entropía se calcula con probabilidades normalizadas dentro de su clase. La media usa T=promedio(I). El método del mínimo selecciona un valle entre dos picos del histograma suavizado; puede fallar si los picos pertenecen al mismo fondo.')
p('La multiumbralización divide la intensidad en tres clases: 0 para I<80, 127 para 80≤I<150 y 255 para I≥150. El umbral por banda produce una máscara blanca únicamente para 80≤I≤150. Las tres clases de intensidad no equivalen a tres objetos.')
h('Ejemplo manual de umbralado')
p('Considérese la matriz I=[[50,50,200],[50,50,200],[200,200,200]]. Hay cuatro píxeles de nivel 50 y cinco de nivel 200. Con T=100 y polaridad inversa, la máscara es [[1,1,0],[1,1,0],[0,0,0]].')
table(['Cantidad','Cálculo','Resultado'], [
    ['Probabilidades','w₀=4/9; w₁=5/9','0.444 y 0.556'],
    ['Varianza intraclase','w₀·0+w₁·0','0'],
    ['Área y componentes','Σ máscara; vecindad 8','4 px; 1 región'],
    ['Perímetro por bordes de celdas','2·(2+2)','8 unidades'],
    ['Perímetro del contorno OpenCV','polígono entre centros de píxel','4 px'],
], [2,2.8,1.7])
p('Cualquier umbral entre 50 y 199 separa estos dos niveles sin varianza interna. El ejemplo también distingue dos convenciones de perímetro. En los resultados de las fotografías se usa exclusivamente arcLength del contorno exterior; no se combina con el perímetro por bordes de celdas.')
h('Ejemplo manual de brillo')
p('Con I=64, r=64/255. La corrección γ=0.5 da s=255√r≈128; γ=2 da s=255r²≈16. Un desplazamiento de +30 produce 94. Estos cálculos muestran que γ<1 aclara y γ>1 oscurece bajo la convención s=255r^γ.')

page('Resultados', 1)
h('Imágenes y región de interés')
p('Se utilizaron las fotografías originales de WhatsApp del 16 de octubre de 2026, a las 12:44:20 PM. La imagen sin sufijo se denominó monedas_separadas y la imagen con sufijo (1), monedas_juntas. Sus tamaños son 1280×667 y 1280×687 píxeles, respectivamente. En ambas se recortó x=490, y=280, ancho=125 y alto=330, con origen en la esquina superior izquierda.')
figure(E/'monedas_separadas'/'escena.png', 'Figura 1. Fotografía con monedas separadas y recorte utilizado. El rectángulo rojo delimita la región de interés.', 6.0)
p('El recorte excluye la cartera, las llaves, el recibo y la botella. Por tanto, el conteo reportado corresponde solo a las monedas dentro de esta región y no al total de objetos presentes en la escena. La segunda fotografía se muestra en los anexos con el mismo procedimiento.')
h('Condiciones comunes')
p('Se trabajó en escala de grises de 8 bits, con conectividad 8 y área mínima de 150 píxeles. Se conservó el conteo bruto y el filtrado por área. Se añadió ruido sintético sal y pimienta a 3 % de las posiciones y ruido gaussiano con media 0 y σ=15 niveles, usando semilla 2026. Los filtros se aplicaron antes del umbralado.')

page('Comparación de métodos de umbralado')
p('La tabla muestra umbral T y componentes con área ≥150 px. O significa imagen original y M significa mediana de 5×5. Antes de estas mediciones no se aplicó relleno ni morfología.')
rows=[]
for label,data in [('Separadas',A),('Juntas',B)]:
    for method in ['Otsu','Kapur','Minimo','Media']:
        ro=threshold_row(data,'Original',method); rm=threshold_row(data,'Mediana 5',method)
        rows.append([label,method,f"{ro['threshold']:.2f}",ro['count'],f"{rm['threshold']:.2f}",rm['count']])
table(['Imagen','Método','T O','N O','T M','N M'],rows,[1.15,1.15,1.05,.9,1.05,1.2])
figure(E/'monedas_separadas'/'umbral_0.png', 'Figura 2. Segmentación sin preprocesamiento en las monedas separadas. El mismo conteo puede acompañar máscaras y áreas muy diferentes.', 4.0)
p('En la imagen original separada, Otsu conservó tres regiones con área total de 5793 px; Kapur, tres con 7181 px; y mínimo, tres con 20111 px. El mínimo incluyó amplias zonas de fondo. La media produjo cuatro regiones. Esta comparación impide considerar el conteo correcto como prueba suficiente de segmentación completa.')

page('Transformaciones de brillo y contraste')
p('Se usó r=I/255 y redondeo al entero más cercano, con saturación en [0,255]. La media μ resume brillo, pero no mide calidad de segmentación. N indica regiones tras Otsu y filtro de área, sin mediana ni morfología.')
formulas = [
('uniforme','Uniforme','CDF con equalizeHist'),
('exponencial','Exponencial','255(1−exp(−r))'),
('rayleigh_guia','Rayleigh de la guía','255√r'),
('hipercubica','Hipercúbica','255r⁴'),
('logaritmica','Logarítmica','255·ln(1+I)/ln(256)'),
('potencia_2','Potencia 2','255r²'),
('gamma_07','Gamma 0.7','255r^0.7'),
('gamma_15','Gamma 1.5','255r^1.5'),
('desplazamiento_30','Desplazamiento +30','clip(I+30)'),
('desplazamiento_m30','Desplazamiento −30','clip(I−30)'),
('contraccion','Contracción','80+120r'),
('expansion','Expansión','255(I−mín)/(máx−mín)'),
]
rows=[]
for key,label,formula in formulas:
    aa=next(v for v in A['brightness'] if v['method']==key)
    bb=next(v for v in B['brightness'] if v['method']==key)
    rows.append([label,formula,aa['mean'],aa['count'],bb['count']])
table(['Técnica','Transformación','μ separadas','N separadas','N juntas'],rows,[1.45,2.1,1.05,.95,.95])
p(f"La media original fue {A['source_mean']:.2f} en las monedas separadas y {B['source_mean']:.2f} en las juntas. Gamma 0.7 aclaró ambas imágenes. La exponencial de la guía oscureció y limitó su salida a un máximo de 161; por su fórmula, no expande todo el rango de 8 bits.")
p('Se conservaron los nombres de la guía para facilitar la comparación. Sin embargo, raíz, potencia, exponencial y logaritmo son transformaciones puntuales de intensidad: no constituyen por sí mismas especificaciones estadísticas de histogramas Rayleigh o exponencial. La ecualización uniforme sí usa la distribución acumulada. Tampoco se debe confundir la función potencia con una técnica distinta de gamma: comparten la misma familia matemática.')

page('Evidencia visual del ajuste de brillo')
figure(E/'monedas_separadas'/'brillo_0.png', 'Figura 3. Uniforme, exponencial, raíz denominada Rayleigh en la guía, hipercúbica, logarítmica y potencia 2 sobre el mismo recorte.', 4.3)
figure(E/'monedas_separadas'/'histogramas.png', 'Figura 4. Histogramas original, ecualizado, con gamma 0.7 y expandido.', 5.8)
p('La uniforme extendió los niveles, pero generó cuatro regiones filtradas en esta imagen. El logaritmo y la raíz aclararon intensamente y también dieron cuatro. Potencia 2 e hipercúbica oscurecieron las monedas y conservaron tres regiones. La mejor apariencia visual no implica automáticamente una máscara más fiel.')

page('Gamma y transformaciones lineales')
figure(E/'monedas_separadas'/'brillo_1.png', 'Figura 5. Gamma 0.7 y 1.5, desplazamientos ±30, contracción a [80,200] y expansión del rango observado.', 4.5)
p('El desplazamiento +30 saturó parte del fondo claro y puede perder diferencias entre intensidades altas. El desplazamiento −30 conservó el conteo de Otsu y trasladó su umbral de 179 a 149. La contracción redujo el rango disponible; la expansión utilizó el mínimo y máximo de cada recorte para redistribuirlo a [0,255].')
p('En la imagen separada, Otsu produjo tres regiones tanto en el original como con gamma 0.7 y gamma 1.5. Después de mediana, gamma 0.7 cambió el umbral de 183 a 200 y el área de la máscara filtrada de 6536 a 6267 px. No se dispone de una máscara de referencia que permita atribuir ese cambio a una mejora de exactitud.')
p('En la imagen con monedas juntas, gamma tampoco eliminó el puente entre las dos monedas superiores. Una transformación global de brillo no incorpora información de forma ni garantiza la separación de objetos que se tocan en la máscara.')

page('Ruido y filtrado antes de segmentar')
p('El ruido se añadió a la ROI original para disponer de una referencia conocida en la comparación. MAE es el error absoluto medio en niveles de gris respecto a esa ROI; no estima el ruido que pudiera existir en la fotografía original.')
rows=[]
for r in A['noise']:
    rows.append([r['noise'],r['filtering'],r['mae'],r['raw_count'],r['count']])
table(['Ruido','Filtro','MAE','N bruto','N ≥150'], rows,[1.45,1.6,1.0,1.15,1.3])
figure(E/'monedas_separadas'/'ruido_filtros.png', 'Figura 6. Ruido sintético y filtrado con ventanas 5×5 sobre las monedas separadas.', 3.7)
p('Ante sal y pimienta, la mediana redujo el conteo bruto de 464 a 4; el gaussiano dejó 108 regiones. Con ruido gaussiano, el filtro gaussiano obtuvo MAE 5.502 y la mediana 5.919; ambos conservaron tres regiones. En este ensayo se favorece mediana para impulsos y gaussiano para ruido gaussiano.')

page('Máscaras finales y medidas')
p('Para comparar una secuencia común se aplicaron mediana 5×5, gamma 0.7, Otsu inverso, cierre elíptico 5×5, relleno de huecos y apertura elíptica 3×3. Se retiraron componentes de área <150 px. El cierre y el relleno completan zonas internas, pero también pueden conectar objetos; sus medidas corresponden a la máscara procesada.')
figure(E/'monedas_separadas'/'final.png', 'Figura 7. Secuencia y contornos finales de las monedas separadas.', 5.7)
figure(E/'monedas_juntas'/'final.png', 'Figura 8. Secuencia idéntica en las monedas juntas. Las dos superiores forman una sola región.', 5.7)
rows=[]
for label,data in [('Separadas',A),('Juntas',B)]:
    for r in data['final']['regions']:
        rows.append([label,r['number'],r['area'],f"{r['perimeter']:.2f}"])
table(['Imagen','Región de arriba abajo','Área px','Perímetro exterior px'],rows,[1.3,1.9,1.3,2.0])
p(f"Totales de la máscara: separadas, {A['final']['count']} regiones, {A['final']['area']} px y {A['final']['perimeter']:.2f} px de perímetro exterior acumulado; juntas, {B['final']['count']} regiones, {B['final']['area']} px y {B['final']['perimeter']:.2f} px. La región 1 de la segunda imagen incluye dos monedas y el puente que las conecta.")

page('Análisis y Discusión', 1)
h('Efecto del preprocesamiento')
p('La mediana redujo las pequeñas regiones causadas por textura o ruido. En la fotografía separada, el conteo bruto de Otsu pasó de 16 en el original a 3 después de mediana; el filtrado por área ya daba tres en ambos casos. En la fotografía junta pasó de 12 a 2. El preprocesamiento estabilizó las máscaras, pero no resolvió la unión de las monedas superiores.')
p('Kapur produjo una región principal en las monedas juntas después de mediana, mientras que Otsu conservó dos. El mínimo fue sensible a los picos elegidos: en las monedas separadas, el umbral 220 incorporó fondo pese a producir tres regiones. La media dependió de la proporción de mesa y objetos dentro del recorte. Estos métodos globales no modelan las diferencias locales de iluminación ni la identidad de cada moneda.')
h('Nitidez y preservación de bordes')
table(['Filtro Aplicado','Tipo de Ruido','Nitidez','Preservación de Bordes','Reducción de Ruido'],[
    ['Mediana 5×5','Sal y pimienta','Suaviza el relieve','Conserva mejor saltos fuertes; altera detalles pequeños','464→4 regiones brutas'],
    ['Gaussiano 5×5','Sal y pimienta','Difumina impulsos','Extiende transiciones','464→108 regiones brutas'],
    ['Mediana 5×5','Gaussiano','Suaviza textura','Puede cambiar bordes finos','MAE 11.774→5.919'],
    ['Gaussiano 5×5','Gaussiano','Desenfoque visible','Redondea transiciones','MAE 11.774→5.502'],
], [1.2,1.1,1.25,1.65,1.3], from_template=True)
p('Las observaciones de nitidez y bordes son cualitativas, basadas en las figuras. No se calculó una métrica de fidelidad del contorno. Un menor MAE indica mayor proximidad en intensidad a la imagen previa al ruido añadido, pero puede coexistir con pérdida de detalle.')
h('Qué se puede concluir')
p('Para las monedas separadas, mediana y Otsu proporcionaron una base útil; el cierre y relleno permitieron obtener regiones más completas. Gamma 0.7 se mantuvo como ensayo de ajuste de brillo, sin evidencia suficiente para declararlo indispensable. Para las monedas juntas, se requiere una etapa que separe formas conectadas, por ejemplo watershed con marcadores. Esa mejora se propone; no se atribuye al experimento ejecutado.')

page('Evaluación Técnica', 1)
h('Implementación y reproducción')
p('El archivo experimento_segmentacion.py usa Python, OpenCV, NumPy y Matplotlib. La aplicación principal mantiene sus funciones de carga, grises, histograma, Otsu, ruido, filtros y componentes. Los métodos adicionales y las figuras de este informe se ejecutan desde el script independiente; no se presentan como funciones ya integradas en los menús de la aplicación.')
table(['Función','Operación y salida'],[
    ['mask','Obtiene T con Otsu, Kapur, mínimo o media y genera blanco para I≤T.'],
    ['kapur','Recorre umbrales, evita clases vacías y maximiza la suma de entropías normalizadas.'],
    ['histogram_minimum','Suaviza h con σ=2; elige dos máximos altos separados ≥20 niveles y busca su valle.'],
    ['transform','Calcula en flotante para evitar desbordamiento; redondea y satura a uint8.'],
    ['measurements','Usa connectedComponentsWithStats; obtiene área por píxeles y perímetro con arcLength.'],
    ['fill_holes','Inunda el fondo desde un borde negro añadido y rellena huecos cerrados.'],
    ['run','Genera PNG, JSON y CSV con parámetros y resultados de ambas fotografías.'],
], [1.6,4.9])
p('Comando desde la raíz del proyecto:')
cmd=p(r'.\env\Scripts\python.exe PracticaSegmentacion/experimento_segmentacion.py')
for run in cmd.runs: run.font.size=Pt(9);run.font.name='Consolas'
p('Las operaciones no lineales se ensayaron en grises. Las constantes, ROI, semilla y área mínima están al inicio del script. La configuración y las versiones utilizadas quedan registradas en evidencia/resultados.json. El archivo comparacion_umbrales.csv conserva las 24 comparaciones de dos imágenes, tres etapas y cuatro umbrales.')
h('Comprobaciones')
p('El script verifica la dirección de gamma, la saturación del desplazamiento, los extremos de la contracción, la separación de un histograma de dos niveles y el área/perímetro de una figura conocida. También se ejecutaron las siete pruebas existentes del proyecto, que pasaron. Estas comprobaciones validan operaciones concretas; no sustituyen una evaluación de exactitud frente a máscaras anotadas.')
p('Kapur maximiza H₀+H₁ con la partición ≤T y >T. Se corrigió el ejemplo del PDF, que pondera las entropías y no mantiene la misma partición al aplicar el umbral.')

page('Conclusiones', 1)
p('La segmentación dependió de la distribución de intensidades y de la separación espacial entre objetos. El conteo final coincidió con las tres monedas visibles cuando estaban separadas, pero produjo dos regiones cuando las dos superiores quedaron unidas. Por tanto, el conteo de componentes debe interpretarse como conteo de regiones conectadas, sujeto a la calidad de la máscara.')
p('La mediana fue útil ante ruido sal y pimienta, mientras que el suavizado gaussiano obtuvo menor MAE ante ruido gaussiano en las monedas separadas. Ambos filtros pueden alterar detalles y bordes. Conviene seleccionar su ventana según el tamaño del objeto y comprobar el efecto antes de medir.')
p('Los ajustes de brillo modificaron el histograma y el umbral seleccionado. Aclarar con gamma, raíz o logaritmo no garantizó una mejor separación; la ecualización uniforme tampoco fue superior en todos los casos. La evaluación debe combinar apariencia visual, continuidad del contorno y rechazo del fondo, además del número de regiones.')
p('Se obtuvieron área y perímetro exterior de las máscaras finales, sin convertirlos a unidades físicas. Estas cantidades cambian con el umbral, la morfología, el encuadre y la distancia de captura. No deben interpretarse como mediciones físicas exactas ni compararse entre fotografías como si tuvieran la misma escala.')
h('Mejoras propuestas y aplicación')
p('Para un sistema de conteo de monedas se propone controlar la iluminación, reducir sombras y mantener la cámara perpendicular a la mesa. También conviene calibrar la escala, registrar una ROI consistente y separar objetos conectados mediante distancia y watershed. Una validación posterior debería usar máscaras trazadas manualmente para medir IoU y error de área, junto con error de conteo; ninguna de esas métricas se atribuye a este ensayo.')
p('La multiumbralización puede ayudar a distinguir materiales o tonalidades, mientras que el umbral por banda permite aislar un intervalo específico. En estas monedas, sus reflejos y aleaciones producen intensidades diferentes dentro de una misma pieza; por ello estas técnicas no identifican automáticamente una moneda completa ni su denominación.')

page('Anexos', 1)
h('Segunda fotografía y segmentación por intervalos')
figure(E/'monedas_juntas'/'escena.png', 'Figura 9. Fotografía con monedas juntas y recorte fijo. El contacto y las sombras favorecen la conexión de las regiones superiores.', 5.5)
figure(E/'monedas_separadas'/'banda_multi.png', 'Figura 10. Tres clases con umbrales 80 y 150 y banda inclusiva [80,150] en las monedas separadas.', 5.4)
p('La banda selecciona parte de la intensidad de las monedas y omite zonas más claras u oscuras de cada pieza. La máscara por banda no se utilizó para los valores finales de área y perímetro.')

page('Orden del ajuste y realce')
figure(E/'monedas_separadas'/'orden.png', 'Figura 11. Gamma antes y después de mediana sobre ruido impulsivo, y realce tras suavizado gaussiano y gamma sobre ruido gaussiano.', 5.2)
p(f"Se comparó mediana(gamma(I)) con gamma(mediana(I)). La diferencia máxima fue {A['order_max_difference']} niveles: las dos secuencias dieron el mismo resultado en esta realización. Esto concuerda con que una transformación monótona y la mediana de orden impar conservan el orden de las muestras. No se generaliza esta equivalencia al filtro gaussiano.")
p('El tercer panel aplica G=gamma(Gaussiano(I)) y después s=clip(G−0.3·Laplaciano(G)). Es un ensayo de realce paso altas posterior al suavizado y al brillo. No se usó como máscara final ni demuestra una mejora cuantitativa del borde. Aplicarlo directamente a imágenes ruidosas puede enfatizar también el ruido.')
h('Resultados de ruido en las monedas juntas')
table(['Ruido','Filtro','MAE','N bruto','N ≥150'],[
    [r['noise'],r['filtering'],r['mae'],r['raw_count'],r['count']] for r in B['noise']
], [1.4,1.6,1.1,1.2,1.2])
p('Con ruido impulsivo, el conteo filtrado de tres regiones sin suavizado es accidental: no confirma que cada moneda esté correctamente delineada. Después de filtrar vuelve a dos, coherente con la conexión observada entre las monedas superiores.')

page('Evidencia adicional y referencias')
figure(E/'monedas_separadas'/'global.png', 'Figura 12. Otsu en la escena completa. La máscara también contiene cartera, llaves, texto y sombras; no equivale a una detección exclusiva de monedas.', 5.8)
h('Archivos de evidencia')
p('La carpeta evidencia contiene las figuras de ambas fotografías, máscaras por método, histogramas, resultados.json, comparacion_umbrales.csv y mascara_final.png para cada caso. El código y las imágenes originales se conservan junto al informe para reproducir los resultados.')
h('Referencias')
p('Cruz Meza, M. E. Práctica 5 Umbralización y ajuste de brillo. Material de la asignatura, PDF proporcionado, pp. 2–4 y 8–14. Plantilla de informe proporcionada por la asignatura.')
p('Kapur, J. N., Sahoo, P. K. y Wong, A. K. C. (1985). A new method for gray-level picture thresholding using the entropy of the histogram. Computer Vision, Graphics, and Image Processing, 29, 273–285. https://doi.org/10.1016/0734-189X(85)90125-2')
p('OpenCV. Structural Analysis and Shape Descriptors. Documentación de connectedComponentsWithStats y arcLength. https://docs.opencv.org/4.13.0/d3/dc0/group__imgproc__shape.html')

# Logos institucionales en el encabezado de la primera página.
section = doc.sections[0]
section.different_first_page_header_footer = True
section.header_distance = Inches(.12)
header = section.first_page_header.paragraphs[0]
header.paragraph_format.space_before = Pt(0)
header.paragraph_format.space_after = Pt(0)
header.paragraph_format.line_spacing = 1
header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
header.add_run().add_picture(str(BASE/'logos'/'ipn-documento.png'), height=Inches(.65))
header.add_run('\t')
header.add_run().add_picture(str(BASE/'logos'/'escom-documento.png'), height=Inches(.65))
for run in header.runs:
    run.font.size = Pt(1)

# Mantener intactas las partes opacas del paquete de referencia.
intermediate = TMP/'informe-generado.docx'
doc.save(intermediate)
editable={'word/document.xml','word/styles.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
with ZipFile(REF) as src, ZipFile(intermediate) as new, ZipFile(OUT,'w',ZIP_DEFLATED) as dest:
    for info in new.infolist():
        if info.filename in src.namelist() and info.filename not in editable:
            dest.writestr(src.getinfo(info.filename),src.read(info.filename))
        else:
            dest.writestr(info,new.read(info.filename))
with ZipFile(REF) as src, ZipFile(OUT) as out:
    for name in src.namelist():
        if name not in editable:
            assert src.read(name)==out.read(name), name
assert Document(OUT).sections[0].page_width == Document(REF).sections[0].page_width
assert hashlib.sha256(REF.read_bytes()).hexdigest() == 'ae4b8d3a92b1d709dd93db71de7090b1f6c9ef45adc0c6aca996c5493f09378a'
print(OUT)
